#!/usr/bin/env python3
"""LẬP KẾ HOẠCH NGUYÊN LIỆU TỪ LỜI ĐỌC -> SHEET ĐỂ ANH DUYỆT.
  python3 plan_assets.py vo/wordts.json --beat B1 --out plan/ [--target 9] [--style-bible "..."]
Đầu vào : wordts.json (Whisper, mốc từng từ) của beat.
Đầu ra  : plan/asset_plan.xlsx  (+ .csv từng sheet, + asset_manifest.json)
  - SHOTS    : mỗi dòng = 1 ảnh/footage cần có (thời gian phủ, lời đọc, mô tả gợi ý, prompt NHÁP, chuyển động đề xuất, tên file đích)
  - OVERLAYS : sự kiện chữ/đồ hoạ theo từng từ khoá (phủ định, số, tên riêng, trích dẫn...)
  - SFX      : âm thanh nền/nhấn theo cảnh
  - TONG_QUAN: số lượng, thời lượng, những gì cần Muse sinh/tìm
Cột DUYỆT có dropdown OK / SỬA / BỎ (mặc định CHỜ). Sau khi anh duyệt chạy tools/apply_approval.py.
LƯU Ý: kế hoạch này là KHUNG theo cấu trúc lời đọc. Mô tả hình/prompt chỉ là NHÁP — Muse viết prompt cuối cùng.
"""
import json, re, os, sys, argparse, csv, unicodedata

ap = argparse.ArgumentParser()
ap.add_argument("wordts"); ap.add_argument("--beat", default="B1"); ap.add_argument("--out", default="plan")
ap.add_argument("--target", type=float, default=9.0, help="độ dài mục tiêu mỗi cảnh (s)")
ap.add_argument("--min", type=float, default=4.0); ap.add_argument("--max", type=float, default=14.0)
ap.add_argument("--gap", type=float, default=0.35, help="khoảng nghỉ tối thiểu tách mệnh đề")
ap.add_argument("--style-bible", default="painterly photorealistic, muted natural palette, soft cinematic light, shallow haze")
ap.add_argument("--preset", default="", help="dossier|newsreel|cinematic|kinetic (ghi chú phong cách dựng)")
ap.add_argument("--fix", default="", help="sửa chính tả Whisper: Kupka=Koepcke,Lancer=LANSA,Pucalpa=Pucallpa")
ap.add_argument("--long-shot", type=float, default=10.0, help="cảnh dài hơn mức này cần thêm ảnh cận/biến thể")
ap.add_argument("--no-rhythm", action="store_true", help="tắt chia độ dài ảnh theo nhịp (mỗi cảnh chia đôi như bản cũ)")
a = ap.parse_args()
os.makedirs(a.out, exist_ok=True)

W = json.load(open(a.wordts, encoding="utf-8"))
FIX = dict(p.split("=", 1) for p in a.fix.split(",") if "=" in p)
def _fix(w):
    b = re.sub(r"[^A-Za-z0-9']", "", w)
    return w.replace(b, FIX[b]) if b in FIX else w
W = [[s_, e_, _fix(w_)] for s_, e_, w_ in W]
norm = lambda s: re.sub(r"[^a-z0-9']", "", s.lower())
toks = [(s, e, w.strip(), norm(w)) for s, e, w in W]
DUR = toks[-1][1]

# ---------------- 1. mệnh đề -> cảnh (cắt trong khoảng lặng)
clauses, cur = [], [toks[0]]
for p, n in zip(toks, toks[1:]):
    if n[0] - p[1] >= a.gap: clauses.append(cur); cur = []
    cur.append(n)
clauses.append(cur)
cl = [(c[0][0], c[-1][1], " ".join(x[2] for x in c)) for c in clauses]
gaps = [cl[i+1][0] - cl[i][1] for i in range(len(cl) - 1)] + [0]

groups, start_i = [], 0
for i in range(len(cl)):
    dur = cl[i][1] - cl[start_i][0]
    nxt = (cl[i+1][1] - cl[start_i][0]) if i + 1 < len(cl) else 0
    cut = (dur >= a.target and gaps[i] >= 0.45) or dur >= a.max or (nxt > a.max and dur >= a.min)
    if cut or i == len(cl) - 1:
        groups.append((start_i, i)); start_i = i + 1
if len(groups) > 1 and (cl[groups[-1][1]][1] - cl[groups[-1][0]][0]) < a.min:
    g2 = groups.pop(); g1 = groups.pop(); groups.append((g1[0], g2[1]))

def cut_point(i):          # điểm bắt đầu dissolve 0,5s nằm trong khoảng lặng sau mệnh đề i
    gs, ge = cl[i][1], cl[i+1][0]
    return round(gs + max(ge - gs - 0.5, 0) / 2, 2)

# ---------------- 2. từ vựng -> thẻ nội dung
LEX = {
 "water":   (["river","stream","water","lake","flood","swamp","pool","drink","thirst","thirsty","wet","creek"],
             "Mặt nước / dòng suối / sông", ["ripple", "parallax", "mist", "particles dust"], ["river-stream", "water ambience"]),
 "forest":  (["forest","jungle","rainforest","trees","canopy","amazon","leaves","vegetation","panguana","woods","station"],
             "Rừng / tán lá / khu trú ẩn trong rừng", ["sway", "parallax", "mist", "particles dust"], ["forest ambience (insects, birds)"]),
 "storm":   (["storm","thunderstorm","thunder","lightning","flash","rain","hurricane","clouds","wind"],
             "Bão / sấm sét / mây dày", ["rain", "glow flicker (tia sét)", "flash + shake"], ["rain heavy", "thunder deep", "wind gusty"]),
 "aircraft":(["plane","flight","airplane","aircraft","airport","takeoff","flew","flying","wing","engine","seat","pilot","crash","crashed","fell","fall","falling","sky","lansa","electra","lockheed"],
             "Máy bay / sân bay / cánh máy bay / bầu trời", ["parallax", "glow", "particles dust"], ["engine/prop distant", "wind gusty"]),
 "dark":    (["night","dark","darkness","silence","alone","shadow","dusk","lost","nothing","silent"],
             "Cảnh tối / cô độc / im lặng", ["particles embers/dust", "glow pulse", "dim (tối dần)"], ["low drone", "near-silence"]),
 "fire":    (["fire","flame","burn","burned","smoke","ember","embers"],
             "Lửa / khói", ["particles embers", "glow flicker", "mist/smoke"], ["fire crackle"]),
 "people":  (["girl","mother","father","family","parents","child","man","woman","villagers","hunters","doctor","rescue","rescuers","woodcutters","loggers","survivor","passengers"],
             "Con người (silhouette / từ phía sau, KHÔNG thấy rõ mặt)", ["parallax nhẹ", "glow nhẹ"], ["room tone / footsteps nhẹ"]),
 "document":(["investigators","report","ruled","licence","license","law","court","evidence","record","official","cause","conditions","wrote","write"],
             "Tài liệu / hồ sơ / báo cũ / bàn làm việc", ["particles dust", "preset newsreel gợi ý"], ["paper/typewriter nhẹ"]),
 "map":     (["lima","pucallpa","pucalpa","route","miles","kilometers","andes","mountains","peru","peruvian","border"],
             "Bản đồ / địa hình / tuyến bay", ["parallax nhẹ", "highlight_box / timeline overlay"], ["wind high altitude"]),
 "home":    (["home","house","christmas","school","graduated","graduation","december"],
             "Nhà / dịp lễ / bối cảnh gia đình", ["glow ấm (đèn)", "particles dust"], ["room tone / ambience ấm"]),
}
def tags_of(text):
    ws = [norm(w) for w in text.split()]
    sc = {k: sum(1 for w in ws if w in v[0]) for k, v in LEX.items()}
    return sorted([k for k in sc if sc[k] > 0], key=lambda k: -sc[k]) or ["generic"]

# ---------------- 2b. NHỊP ẢNH: độ dài từng ảnh KHÁC NHAU theo loại cảnh (chống slideshow đều tay)
# vai trò: wide=toàn cảnh (ảnh mới) · mid=trung cảnh/góc khác (ảnh mới) · detail=CẮT SÂU vào chi tiết của ảnh trước (KHÔNG cần ảnh mới)
#          hold=giữ hình tự nói (ảnh mới, không chữ/hiệu ứng, spec "hold": true)
PAT = {
 "dark":     [("hold", 7.0), ("detail", 2.5), ("mid", 4.0)],
 "storm":    [("wide", 3.5), ("detail", 2.2), ("mid", 3.0), ("detail", 2.2)],
 "aircraft": [("wide", 3.5), ("detail", 2.2), ("mid", 3.0), ("detail", 2.2)],
 "fire":     [("wide", 3.5), ("detail", 2.2), ("mid", 3.0)],
 "people":   [("wide", 5.0), ("detail", 3.0), ("mid", 6.0)],
 "home":     [("wide", 5.0), ("detail", 3.0), ("mid", 6.0)],
 "document": [("wide", 5.0), ("detail", 3.0), ("mid", 5.0)],
 "map":      [("wide", 5.0), ("detail", 3.0), ("mid", 5.0)],
 "forest":   [("wide", 6.0), ("detail", 2.8), ("mid", 4.5)],
 "water":    [("wide", 6.0), ("detail", 2.8), ("mid", 4.5)],
 "_":        [("wide", 5.0), ("detail", 3.0), ("mid", 5.0)],
}
def plan_shots(t0, t1, cands, main):
    L = t1 - t0
    if a.no_rhythm:
        return [(t0, t1, "wide")] if L <= a.long_shot else [(t0, t0 + L/2, "wide"), (t0 + L/2, t1, "mid")]
    if L <= 6.0: return [(t0, t1, "hold" if main == "dark" and L >= 4 else "wide")]
    pat = PAT.get(main, PAT["_"]); out = []; cur = t0; k = 0
    while cur < t1 - 0.01:
        role, tl = pat[k % len(pat)]; k += 1
        want = cur + tl
        if t1 - want < 2.0: nxt = t1
        else:
            c = [x for x in cands if x > cur + 1.6 and x < t1 - 1.5]
            nxt = min(c, key=lambda x: abs(x - want)) if c else t1
            if t1 - nxt < 2.0: nxt = t1
        out.append((cur, nxt, role)); cur = nxt
    return out
DESC = {"wide": "toàn cảnh (wide)", "mid": "trung cảnh / góc khác, CÙNG bối cảnh",
        "hold": "toàn cảnh tĩnh, HOLD để hình tự nói (không chữ, không hiệu ứng)"}

# ---------------- 3. sự kiện overlay theo từ khoá
NEG = {"not","no","never","nothing","none","nobody","neither","nor","didn't","wasn't","couldn't","wouldn't","isn't","don't","cannot","can't"}
NUMW = {"one","two","three","four","five","six","seven","eight","nine","ten","eleven","twelve","twenty","thirty","forty","fifty","hundred","thousand","million","half","first","second","third"}
TIMEW = {"days","day","hours","minutes","weeks","months","years","later","earlier","before","after"}
events = []
for c in clauses:
    text = " ".join(x[2] for x in c)
    for k, (s, e, w, n) in enumerate(c):
        wl = w.lower().strip(".,;:?!\"“”")
        if n in NEG or n.endswith("nt"):
            events.append((s, w, "PHỦ ĐỊNH", "strike_cards", f"thẻ ghi đối tượng bị phủ định trong: “{text}”"))
        elif re.fullmatch(r"\d[\d,]*", n) or n in NUMW:
            events.append((s, w, "CON SỐ", "ticker / badge_pop", "số chạy tới đích đúng lúc nói"))
        elif n in TIMEW:
            events.append((s, w, "MỐC THỜI GIAN", "timeline / label / stamp", "ngày/mốc thời gian"))
        elif w[:1].isupper() and k > 0 and wl not in ("i",):
            events.append((s, w, "TÊN RIÊNG", "lower_third / highlight_box", f"nhãn tên: {wl.upper()}"))
    if text.rstrip().endswith("?"): events.append((c[0][0], text.split()[0], "CÂU HỎI", "question_card", "dấu ? nhấp nháy"))
    if re.search(r"\b(said|words were|wrote|ruled|called)\b", text, re.I):
        events.append((c[0][0], text.split()[0], "TRÍCH DẪN/PHÁN QUYẾT", "label dẫn + typed_card", "nhãn dẫn rồi gõ chữ thường"))
events.sort()

# ---------------- 3b. QUÉT ÂM THANH THEO TỪ KHOÁ (tiếng một lần / hành động) — ngoài nền theo cảnh
SFXKW = [
 (("crash","crashed","hit","broke","broken","tore","burst","snapped","smashed","slammed","landed","struck","collided","exploded"), "va chạm / vỡ / rơi đập (impact)", "impact_hit"),
 (("walked","walking","steps","footsteps","stepped","ran","running","waded","trudged","stumbled"), "bước chân (đất ẩm / lá / nước nông)", "footsteps_soft"),
 (("splash","splashed","swam","waded","drowned"), "nước té / lội nước", "water_wade"),
 (("vulture","bird","birds","owl","parrot","monkey","howler","jaguar","roar","growl","howl","barked","bark"), "tiếng thú / chim (xa)", "animal_call"),
 (("screamed","shouted","cried","whispered","spoke","called","yelled","voices","voice","sobbed","prayed"), "tiếng người xa / thì thầm / kêu", "voice_distant"),
 (("door","hut","boat","canoe","motor","engine","outboard","propeller"), "cửa kẹt / thuyền / động cơ nhỏ", "boat_motor_door"),
 (("heart","heartbeat","breath","breathing","gasped","panting"), "nhịp tim / thở", "heartbeat_breath"),
 (("planes","aircraft","helicopter","helicopters","airplane","overhead"), "máy bay bay qua ở xa", "plane_distant_overhead"),
 (("bell","bells","church","radio","telephone","phone","typewriter"), "chuông / radio / điện thoại", "bell_radio"),
 (("paper","pages","letter","file","document","documents"), "giấy / lật hồ sơ", "paper_rustle"),
 (("fire","burning","crackling","flames"), "lửa lách tách", "fire_crackle"),
]
SFX_ONE = {}      # file -> [(t, từ)]
SFX_DESC = {}
for k_, (words_, desc_, fn_) in enumerate(SFXKW):
    for (st, en, w, n) in toks:
        if n in words_:
            SFX_ONE.setdefault(fn_, []).append((st, w)); SFX_DESC[fn_] = desc_

# ---------------- 4. dựng bảng
SHOT_PLAN = []
rows_shots, rows_sfx, manifest = [], [], {"beat": a.beat, "duration": DUR, "preset": a.preset, "items": []}
scene_info = []
for gi, (i0, i1) in enumerate(groups, 1):
    t0 = 0.0 if gi == 1 else cut_point(groups[gi-2][1])
    t1 = DUR if gi == len(groups) else cut_point(i1) + 0.5
    text = " ".join(cl[i][2] for i in range(i0, i1 + 1))
    tg = tags_of(text); main = tg[0]
    label, fx, sfx = (LEX[main][1], list(LEX[main][2]), list(LEX[main][3])) if main in LEX else ("Cảnh bối cảnh chung / ẩn dụ", ["parallax nhẹ", "particles dust"], ["room tone"])
    for t2 in tg[1:3]:     # thẻ phụ cũng đóng góp chuyển động + âm thanh (vd: cảnh máy bay có bão -> mưa, sấm)
        if t2 in LEX and t2 not in ("people", "home"):
            fx += [x for x in LEX[t2][2] if x not in fx]; sfx += [x for x in LEX[t2][3] if x not in sfx]
    mood = "căng thẳng, tối" if main in ("storm", "dark") else ("ấm, hoài niệm" if main in ("home", "document") else "tự nhiên, thoáng")
    L = t1 - t0
    cands = []                      # điểm cắt khả dĩ: mọi chỗ ngắt hơi ≥0.2s trong cảnh (dấu phẩy cũng tính)
    for (p_, n_) in zip(toks, toks[1:]):
        if n_[0] - p_[1] >= 0.2 and t0 + 1.5 < (p_[1] + n_[0]) / 2 < t1 - 1.5: cands.append(round((p_[1] + n_[0]) / 2, 2))
    plan = plan_shots(t0, t1, cands, main)
    SHOT_PLAN.append((gi, plan, main))
    scene_info.append((gi, t0, t1, main, tg))
    # nhóm: mỗi ảnh MỚI (wide/mid/hold) + các cú cắt sâu (detail) đi kèm ngay sau nó
    imgs = []
    for (ts, te, role) in plan:
        lim = 9.0 if role == "detail" else 8.0      # "mid" ngắn cũng chỉ là cú cắt sâu/đổi khung từ ảnh có sẵn; dài mới cần ảnh mới
        if role in ("detail", "mid") and imgs and (te - imgs[-1]["start"]) <= lim and (te - ts) <= 4.5:
            imgs[-1]["end"] = te; imgs[-1]["cuts"].append((ts, te)); continue
        imgs.append({"start": ts, "end": te, "role": "mid" if role == "detail" else role, "cuts": []})
    for vi, im_ in enumerate(imgs):
        v = chr(65 + vi); sid = f"{a.beat}-S{gi:02d}{v}"
        fname = f"stills/{a.beat}_S{gi:02d}{v}.png"
        ts, te = im_["start"], im_["end"]; role = im_["role"]
        covers = f"{ts:.1f}–{te:.1f}s"
        stext = " ".join(cl[i][2] for i in range(i0, i1 + 1) if cl[i][1] > ts and cl[i][0] < te) or text
        vdesc = DESC.get(role, DESC["mid"])
        prompt = (f"{label} ({vdesc}). Gắn với lời: “{stext[:120]}…”. Không khí: {mood}. Phong cách: {a.style_bible}. "
                  f"16:9, 2560x1440, KHÔNG chữ/logo, không thấy rõ khuôn mặt, khung hình còn chỗ trống cho chữ.")
        mv = list(fx)
        if role == "hold": mv = ["HOLD: giữ yên, camera trôi rất chậm; KHÔNG chữ/hiệu ứng (spec: scene \"hold\": true)"] + mv[:2]
        for (c0, c1) in im_["cuts"]: mv.append(f"CẮT SÂU vào chi tiết {c0:.1f}–{c1:.1f}s (zoom từ chính ảnh này, KHÔNG cần ảnh mới)")
        rows_shots.append([sid, gi, covers, round(te - ts, 1), stext, f"{label} — {vdesc}",
                           "AI ảnh 2560x1440 (bắt buộc) · footage/AI video 6–10s (tuỳ chọn, nên có cho: " + ("nước/mưa/lửa" if main in ('water','storm','fire') else "cảnh đang hành động") + ")",
                           prompt, ", ".join(mv), fname, "CHỜ", ""])
        manifest["items"].append({"id": sid, "kind": "still", "file": fname, "covers": [round(ts, 2), round(te, 2)], "scene": gi, "role": role, "approval": "CHỜ"})
    for s in sfx:
        ascii_ = unicodedata.normalize("NFKD", s.split("(")[0]).encode("ascii", "ignore").decode().lower()
        fn = f"sfx/{re.sub('[^a-z0-9]+', '_', ascii_).strip('_')}.wav"
        prev = next((r for r in rows_sfx if r[5] == fn), None)
        if prev:                                   # cùng loại âm thanh ở nhiều cảnh -> gộp 1 dòng
            prev[1] += f", {gi}"; prev[2] += f"; {t0:.1f}–{t1:.1f}s"
            continue
        sid = f"{a.beat}-SFX{len(rows_sfx)+1:02d}"
        rows_sfx.append([sid, str(gi), f"{t0:.1f}–{t1:.1f}s", "nền (bed)", s, fn, "CHỜ", ""])
for r in rows_sfx:
    manifest["items"].append({"id": r[0], "kind": "sfx", "file": r[5], "when": r[2], "approval": "CHỜ"})
rows_sfx.append([f"{a.beat}-MUS-01", "—", f"0–{DUR:.0f}s", "nhạc nền/drone", "drone/nốt nhạc căng thẳng, rất nhỏ (−25 dB), vào chậm, tắt trước câu chốt", "sfx/music_bed.wav", "CHỜ", ""])
manifest["items"].append({"id": f"{a.beat}-MUS-01", "kind": "sfx", "file": "sfx/music_bed.wav", "approval": "CHỜ"})
ACC = [e for e in events if e[2] in ("CON SỐ", "PHỦ ĐỊNH", "TRÍCH DẪN/PHÁN QUYẾT")][:8]
rows_sfx.append([f"{a.beat}-ACC-01", "—", ", ".join(f"{e[0]:.1f}s" for e in ACC), "nhấn (accent hit)", "1 tiếng nhấn ngắn dùng lại ở các mốc số/phủ định/chốt (không dùng bản tổng hợp tạm nếu có thể)", "sfx/accent_hit.wav", "CHỜ", ""])
manifest["items"].append({"id": f"{a.beat}-ACC-01", "kind": "sfx", "file": "sfx/accent_hit.wav", "approval": "CHỜ"})

# ---- nhịp HOLD: video >30s phải có ít nhất 1 cảnh để hình tự nói
if DUR > 30 and not any(it.get("role") == "hold" for it in manifest["items"]):
    best = None
    for r_, it in zip(rows_shots, [i for i in manifest["items"] if i["kind"] == "still"]):
        ts_, te_ = it["covers"]
        if te_ - ts_ < 5: continue
        n_ev = sum(1 for e in events if ts_ <= e[0] < te_)
        if best is None or n_ev < best[0]: best = (n_ev, r_, it)
    if best:
        best[1][8] = "HOLD: giữ yên, camera trôi rất chậm; KHÔNG chữ/hiệu ứng (spec: scene \"hold\": true) · " + best[1][8]
        best[1][5] += " — HOLD để hình tự nói"; best[2]["role"] = "hold"

# ---- âm thanh một lần theo từ khoá (hành động/vật thể) — gộp theo file
for fn_, hits in SFX_ONE.items():
    hits = sorted(hits)
    ts_ = ", ".join(f"{t:.1f}s('{w.strip('.,;?!')}')" for t, w in hits[:6])
    sid = f"{a.beat}-ONE{len([r for r in rows_sfx if '-ONE' in r[0]])+1:02d}"
    rows_sfx.append([sid, "—", ts_, "hiệu ứng một lần", SFX_DESC[fn_] + " — đặt ĐÚNG lúc từ khoá, ngắn, nhỏ; không lặp quá 3–4 lần", f"sfx/{fn_}.wav", "CHỜ", ""])
    manifest["items"].append({"id": sid, "kind": "sfx", "file": f"sfx/{fn_}.wav", "approval": "CHỜ"})

# ---- ÂM DẪN: tiếng của cảnh sau vào trước khi đổi hình (0.8s)
for (g0, t0_, t1_, m0, tg0), (g1, u0, u1, m1, tg1) in zip(scene_info, scene_info[1:]):
    if m1 != m0 and m1 in LEX:
        nm = LEX[m1][3][0]
        fn_ = "sfx/" + re.sub("[^a-z0-9]+", "_", unicodedata.normalize("NFKD", nm.split("(")[0]).encode("ascii", "ignore").decode().lower()).strip("_") + ".wav"
        rows_sfx.append([f"{a.beat}-LEAD{g1:02d}", f"{g0}→{g1}", f"{u0-0.8:.1f}–{u0+0.5:.1f}s", "âm dẫn (vào trước khi đổi hình)",
                         f"tiếng '{nm}' của cảnh {g1} vào nhỏ trước điểm cắt {u0:.1f}s ~0.8s, rồi lớn dần khi hình đổi (DÙNG LẠI file nền, không cần file mới)", fn_, "CHỜ", ""])

rows_ov = []
for k, (s, w, kind, eff, note) in enumerate(events, 1):
    sc_ = next((g for g, t0, t1, m, tg in scene_info if t0 <= s < t1 + 0.01), "?")
    rows_ov.append([f"{a.beat}-EV{k:03d}", sc_, round(s, 2), w, kind, eff, note, "CHỜ", ""])

# ---------------- 5. ghi ra
H_SHOT = ["ID", "CẢNH", "PHỦ THỜI GIAN", "DÀI (s)", "LỜI ĐỌC", "HÌNH GỢI Ý", "LOẠI NGUYÊN LIỆU", "PROMPT NHÁP (Muse viết bản cuối)", "CHUYỂN ĐỘNG ĐỀ XUẤT", "TÊN FILE ĐÍCH", "DUYỆT", "GHI CHÚ CỦA ANH"]
H_OV = ["ID", "CẢNH", "GIÂY", "TỪ KÍCH HOẠT", "LOẠI", "HIỆU ỨNG ĐỀ XUẤT", "GHI CHÚ NỘI DUNG", "DUYỆT", "GHI CHÚ CỦA ANH"]
H_SFX = ["ID", "CẢNH", "THỜI GIAN", "VAI TRÒ", "MÔ TẢ", "TÊN FILE ĐÍCH", "DUYỆT", "GHI CHÚ CỦA ANH"]
for name, hdr, rows in (("shots", H_SHOT, rows_shots), ("overlays", H_OV, rows_ov), ("sfx", H_SFX, rows_sfx)):
    with open(os.path.join(a.out, f"{name}.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(hdr); w.writerows(rows)
json.dump(manifest, open(os.path.join(a.out, "asset_manifest.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
wb = Workbook()
ws0 = wb.active; ws0.title = "HUONG_DAN"
guide = [
 f"KẾ HOẠCH NGUYÊN LIỆU — BEAT {a.beat}  ({DUR:.0f}s lời đọc, {len(groups)} cảnh)",
 "",
 "Sheet SHOTS = MAPPING câu theo voice (bước 3) + VISUAL PLAN (bước 4): Muse sửa cột HÌNH GỢI Ý thành ý tưởng hình cuối cho từng dòng trước khi gửi anh.",
 "CÁCH DUYỆT (bước 5): ở cột DUYỆT chọn OK (làm đúng như đề xuất), SỬA (làm nhưng theo ghi chú của anh) hoặc BỎ (không làm); xác nhận hoặc đổi style ở dòng cuối sheet này.",
 "Dòng để CHỜ = chưa duyệt, Muse KHÔNG được sinh/tìm nguyên liệu cho dòng đó. Ghi yêu cầu riêng vào cột GHI CHÚ CỦA ANH.",
 "Prompt trong sheet SHOTS chỉ là NHÁP; Muse viết prompt cuối rồi sinh ảnh. Đây là lần duyệt duy nhất trước khi tạo nguyên liệu; sau đó anh xem video chương.",
 "Sau khi duyệt xong: lưu file, đưa lại cho Muse (hoặc đưa Claude chạy tools/apply_approval.py) để ra MUSE-TODO.md.",
 "",
 f"STYLE ĐỀ XUẤT (STYLES.md) — style dựng: {a.preset or 'mặc định (hồ sơ điều tra)'} · style hình: {a.style_bible}",
 "STYLE ANH CHỐT (để trống = đồng ý đề xuất): ",
]
for i, g in enumerate(guide, 1):
    ws0.cell(i, 1, g).alignment = Alignment(wrap_text=True, vertical="top")
ws0["A1"].font = Font(bold=True, size=14); ws0.column_dimensions["A"].width = 120

hf = PatternFill("solid", fgColor="1F3A2E"); thin = Side(style="thin", color="BBBBBB")
def sheet(title, hdr, rows, widths, approve_col):
    ws = wb.create_sheet(title)
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hf; c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(r)
    for i, wd in enumerate(widths, 1): ws.column_dimensions[chr(64 + i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.freeze_panes = "B2"
    col = chr(64 + approve_col)
    dv = DataValidation(type="list", formula1='"CHỜ,OK,SỬA,BỎ"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{len(rows)+1}")
    rng = f"{col}2:{col}{len(rows)+1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SỬA"'], fill=PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"BỎ"'], fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")))
    return ws
sheet("SHOTS", H_SHOT, rows_shots, [14, 7, 13, 8, 46, 32, 36, 70, 28, 24, 9, 30], 11)
sheet("OVERLAYS", H_OV, rows_ov, [14, 7, 8, 16, 20, 26, 48, 9, 30], 8)
sheet("SFX", H_SFX, rows_sfx, [16, 7, 16, 16, 56, 30, 9, 30], 7)
ws = wb.create_sheet("TONG_QUAN", 1)
n_still = len(rows_shots)
summ = [["Chỉ số", "Giá trị"],
        ["Thời lượng lời đọc", f"{DUR:.1f}s"], ["Số cảnh", len(groups)], ["Số ảnh AI cần sinh (bắt buộc)", n_still],
        ["Footage / AI video (tuỳ chọn)", sum(1 for _, _, _, m, _ in scene_info if m in ("water", "storm", "fire", "forest"))],
        ["Số sự kiện overlay đề xuất", len(rows_ov)],
        ["Số cú CẮT SÂU từ ảnh có sẵn (không cần ảnh mới)", sum(len(re.findall("CẮT SÂU", r[8])) for r in rows_shots)],
        ["Số nhịp HOLD (hình tự nói)", sum(1 for i_ in manifest["items"] if i_.get("role") == "hold")],
        ["Độ dài ảnh ngắn nhất / dài nhất", f"{min(r[3] for r in rows_shots):.1f}s / {max(r[3] for r in rows_shots):.1f}s"], ["Số âm thanh cần có", len(rows_sfx)],
        ["Mật độ ảnh", f"1 ảnh / {DUR/n_still:.1f}s"], ["", ""], ["Cảnh", "Thời gian · thẻ nội dung"]]
for g, t0, t1, m, tg in scene_info: summ.append([f"S{g:02d}", f"{t0:.1f}–{t1:.1f}s · {', '.join(tg[:3])}"])
for r in summ: ws.append(r)
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 60
wb.save(os.path.join(a.out, "asset_plan.xlsx"))
print(f"Beat {a.beat}: {len(groups)} cảnh, {n_still} ảnh, {len(rows_ov)} sự kiện overlay, {len(rows_sfx)} âm thanh -> {a.out}/asset_plan.xlsx")
for g, t0, t1, m, tg in scene_info: print(f"  S{g:02d} {t0:6.1f}-{t1:6.1f}s  {','.join(tg[:3])}")
