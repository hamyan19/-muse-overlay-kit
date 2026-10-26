#!/usr/bin/env python3
"""Thêm 3 cột CỤM Ý vào sheet SHOTS của asset_plan.xlsx có sẵn (bảng cũ chưa có cột): CỤM · VÙNG · CHUYỂN TRONG CỤM.

  python3 tools/add_cluster_cols.py plan/CH3_asset_plan.xlsx              # thêm cột trống (ghi đè file, có bản .bak)
  python3 tools/add_cluster_cols.py plan/CH3_asset_plan.xlsx --set "CH3-M01:S04-S07" "CH3-M02:S10-S12:ô"
      --set <mã cụm>:<shot đầu>-<shot cuối>[:ô]   gán nhanh một dải shot vào cụm; ':ô' = lộ dần từng ô (mặc định zoom từng vùng)
Luật gom/tách cụm: KY-THUAT-DUNG.md mục 1. VÙNG: mô tả + (tuỳ chọn) toạ độ '480,820 x2' trên ảnh nguồn 2560x1440.
"""
import argparse, re, shutil, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLS = ["CỤM", "VÙNG", "CHUYỂN TRONG CỤM"]
CHOICES = "mở toàn cảnh,zoom vào,lia,cắt cận,lùi toàn cảnh,mở ô"
ap = argparse.ArgumentParser(); ap.add_argument("xlsx"); ap.add_argument("--set", nargs="*", default=[])
a = ap.parse_args()
shutil.copy(a.xlsx, a.xlsx + ".bak")
wb = openpyxl.load_workbook(a.xlsx); ws = wb["SHOTS"]
H = [c.value for c in ws[1]]
for name in COLS:
    if name not in H:
        c = ws.cell(1, len(H) + 1, name); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A2E")
        ws.column_dimensions[c.column_letter].width = 16 if name == "CỤM" else 34
        H.append(name)
ci = {h: i + 1 for i, h in enumerate(H)}
L = ws.cell(1, ci["CHUYỂN TRONG CỤM"]).column_letter
dv = DataValidation(type="list", formula1=f'"{CHOICES}"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{L}2:{L}{ws.max_row}")
ids = {str(ws.cell(r, ci["ID"]).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, ci["ID"]).value}
for spec in a.set:
    m = re.match(r"([^:]+):S?(\d+)-S?(\d+)(:ô)?$", spec.strip())
    if not m: sys.exit("sai cú pháp --set: " + spec)
    cum, s0, s1, panel = m.group(1), int(m.group(2)), int(m.group(3)), bool(m.group(4))
    tag = cum.split("-")[0]
    for n in range(s0, s1 + 1):
        r = ids.get(f"{tag}-S{n:02d}")
        if not r: print("không thấy shot", f"{tag}-S{n:02d}"); continue
        ws.cell(r, ci["CỤM"], cum)
        if not ws.cell(r, ci["CHUYỂN TRONG CỤM"]).value:
            ws.cell(r, ci["CHUYỂN TRONG CỤM"], "mở ô" if panel else ("mở toàn cảnh" if n == s0 else "zoom vào"))
        if not ws.cell(r, ci["VÙNG"]).value:
            ws.cell(r, ci["VÙNG"], f"ô {n - s0 + 1}" if panel else ("toàn cảnh" if n == s0 else "(mô tả chi tiết + vị trí)"))
for row in ws.iter_rows(min_row=2):
    for c in row[-3:]: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(a.xlsx)
print(f"{a.xlsx}: đã có cột {', '.join(COLS)}" + (f"; gán {len(a.set)} cụm" if a.set else ""))
