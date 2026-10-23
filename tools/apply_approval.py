#!/usr/bin/env python3
"""ĐỌC SHEET ĐÃ DUYỆT -> MUSE-TODO.md + approved_manifest.json
  python3 apply_approval.py plan/asset_plan.xlsx --out plan/
Quy tắc: OK = làm theo đề xuất · SỬA = làm theo GHI CHÚ của anh · BỎ = không làm · CHỜ = CHƯA duyệt, KHÔNG làm.
Nếu còn dòng CHỜ ở sheet SHOTS, script báo rõ và thoát mã 2 (Muse không được tự coi là đã duyệt).
"""
import sys, json, os, argparse, re
import openpyxl
ap = argparse.ArgumentParser(); ap.add_argument("xlsx"); ap.add_argument("--out", default=".")
a = ap.parse_args()
wb = openpyxl.load_workbook(a.xlsx, data_only=True)
def rows(name):
    ws = wb[name]; hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]: yield dict(zip(hdr, r))
out = {"OK": [], "SỬA": [], "BỎ": [], "CHỜ": []}
items = []
for sh in ("SHOTS", "OVERLAYS", "SFX"):
    for r in rows(sh):
        st = (r.get("DUYỆT") or "CHỜ").strip().upper()
        st = {"OK": "OK", "SỬA": "SỬA", "SUA": "SỬA", "BỎ": "BỎ", "BO": "BỎ"}.get(st, "CHỜ")
        it = {"sheet": sh, "id": r["ID"], "status": st, "note": r.get("GHI CHÚ CỦA ANH") or "",
              "file": r.get("TÊN FILE ĐÍCH"), "covers": r.get("PHỦ THỜI GIAN") or r.get("THỜI GIAN") or r.get("GIÂY"),
              "prompt_draft": r.get("PROMPT NHÁP (Muse viết bản cuối)"), "effect": r.get("HIỆU ỨNG ĐỀ XUẤT"), "desc": r.get("MÔ TẢ") or r.get("HÌNH GỢI Ý")}
        txt = it["covers"]
        nums = re.findall(r"\d+(?:\.\d+)?", str(txt or "")) if sh == "SHOTS" else []
        it["covers_text"] = txt
        it["covers"] = [float(nums[0]), float(nums[1])] if len(nums) >= 2 else txt
        if sh == "SHOTS": it["kind"] = "still"
        out[st].append(it); items.append(it)
json.dump({"items": items}, open(os.path.join(a.out, "approved_manifest.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
L = [f"# MUSE-TODO (từ {os.path.basename(a.xlsx)})", "",
     f"Tổng: OK {len(out['OK'])} · SỬA {len(out['SỬA'])} · BỎ {len(out['BỎ'])} · CHƯA DUYỆT {len(out['CHỜ'])}", ""]
def sect(title, lst, show_note=False):
    if not lst: return
    L.append(f"## {title}")
    for it in lst:
        L.append(f"- **{it['id']}** ({it['sheet']}, {it['covers_text']}) → `{it['file']}`" if it['file'] else f"- **{it['id']}** ({it['sheet']}, {it['covers_text']}): {it['effect']}")
        if show_note and it["note"]: L.append(f"  - YÊU CẦU CỦA ANH: {it['note']}")
        elif it["note"]: L.append(f"  - ghi chú: {it['note']}")
    L.append("")
sect("LÀM NGAY (đã duyệt OK)", out["OK"])
sect("LÀM THEO YÊU CẦU RIÊNG CỦA ANH (SỬA) — ghi chú của anh thắng đề xuất ban đầu", out["SỬA"], True)
sect("KHÔNG LÀM (BỎ)", out["BỎ"])
if out["CHỜ"]:
    L.append("## CHƯA DUYỆT — KHÔNG ĐƯỢC LÀM, KHÔNG ĐƯỢC TỰ COI LÀ ĐÃ DUYỆT")
    L += [f"- {it['id']} ({it['sheet']})" for it in out["CHỜ"]]; L.append("")
L += ["## Bước tiếp theo (QUY-TRINH.md bước 6 → 7)", "1. Sinh/tìm nguyên liệu cho OK + SỬA (prompt cuối do Muse viết, theo tên file đích).",
      "2. Một lần kiểm theo QC-ASSETS.md: `python3 tools/qc_assets.py --root res --manifest plan/approved_manifest.json --out qc/` (FAIL=0) + 5 mục mắt.",
      "3. Giao gói cho người dựng: spec → lint → render → gửi anh xem video chương."]
open(os.path.join(a.out, "MUSE-TODO.md"), "w", encoding="utf-8").write("\n".join(L))
print("\n".join(L[:4])); 
pend = [i for i in out["CHỜ"] if i["sheet"] == "SHOTS"]
if pend:
    print(f"\n!! CÒN {len(pend)} DÒNG ẢNH CHƯA DUYỆT — chưa được bắt đầu sinh ảnh cho các dòng này"); sys.exit(2)
